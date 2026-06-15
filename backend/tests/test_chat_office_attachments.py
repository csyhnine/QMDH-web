import io
import unittest

from docx import Document
from openpyxl import Workbook

from app.services.chat_file_text import extract_attachment_text
from app.services.media_storage import write_binary_asset


def _build_docx_bytes(*paragraphs: str) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_xlsx_bytes(rows: list[tuple[str, ...]], *, sheet_title: str = "Sheet1") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class ChatOfficeAttachmentTests(unittest.TestCase):
    def test_extract_docx_text(self) -> None:
        file_path = write_binary_asset(
            "chat-attachments/brief.docx",
            _build_docx_bytes("项目背景", "需要优化主视觉"),
        )
        extracted = extract_attachment_text(file_path, file_name="brief.docx")
        self.assertIn("项目背景", extracted)
        self.assertIn("需要优化主视觉", extracted)

    def test_extract_xlsx_text(self) -> None:
        file_path = write_binary_asset(
            "chat-attachments/budget.xlsx",
            _build_xlsx_bytes([("科目", "金额"), ("印刷", "1200")]),
        )
        extracted = extract_attachment_text(file_path, file_name="budget.xlsx")
        self.assertIn("[Sheet1]", extracted)
        self.assertIn("科目", extracted)
        self.assertIn("印刷", extracted)
        self.assertIn("1200", extracted)


if __name__ == "__main__":
    unittest.main()
